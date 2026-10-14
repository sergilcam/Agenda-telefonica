"""Agenda telefonica de consola.

Proyecto educativo de 2.o de Bachillerato.
Usa programacion estructurada: solo funciones y una lista de diccionarios.
Cada contacto es un diccionario: {"nombre": ..., "telefono": ...}.
El telefono se guarda como texto para conservar ceros y el prefijo +34.
Los datos se guardan en una hoja de calculo Excel (.xlsx)
con dos columnas: nombre | telefono.
"""

import openpyxl


RUTA_DATOS = "data/contactos.xlsx"
NOMBRE_HOJA = "Contactos"
TITULO_NOMBRE = "nombre"
TITULO_TELEFONO = "telefono"


def mostrar_menu():
    """Muestra por pantalla las opciones disponibles."""
    print("\n--- AGENDA TELEFONICA ---")
    print("1. Anadir contacto")
    print("2. Buscar por nombre")
    print("3. Buscar por numero")
    print("4. Eliminar contacto")
    print("5. Mostrar todos los contactos")
    print("0. Salir")


def pedir_opcion():
    """Pide la opcion del menu y la devuelve sin espacios."""
    opcion = input("Elige una opcion: ")
    return opcion.strip()


def pedir_nombre():
    """Pide un nombre y lo devuelve sin espacios delante ni detras."""
    nombre = input("Nombre: ")
    return nombre.strip()


def pedir_telefono():
    """Pide un telefono y lo devuelve sin espacios delante ni detras."""
    telefono = input("Telefono: ")
    return telefono.strip()


def pedir_orden():
    """Pregunta como ordenar y devuelve True si es descendente.

    1 = ascendente (A-Z), 2 = descendente (Z-A).
    Si se pulsa Enter o algo no valido, usa ascendente.
    """
    print("Orden: 1. Ascendente (A-Z)  2. Descendente (Z-A)")
    respuesta = input("Elige el orden [1]: ")
    respuesta = respuesta.strip()
    if respuesta == "2":
        return True
    return False


def crear_contacto(nombre, telefono):
    """Crea un contacto con un nombre y un telefono."""
    return {"nombre": nombre, "telefono": telefono}


def anadir_contacto(agenda, contacto):
    """Anade un contacto a la lista agenda."""
    agenda.append(contacto)


def buscar_por_nombre(agenda, nombre):
    """Busca contactos cuyo nombre contenga el texto dado.

    No distingue entre mayusculas y minusculas.
    Devuelve una lista con los encontrados (puede estar vacia).
    """
    encontrados = []
    texto = nombre.strip().lower()
    for contacto in agenda:
        if texto in contacto["nombre"].lower():
            encontrados.append(contacto)
    return encontrados


def buscar_por_numero(agenda, numero):
    """Busca contactos cuyo telefono contenga el texto dado.

    El numero se compara como texto.
    Devuelve una lista con los encontrados (puede estar vacia).
    """
    encontrados = []
    texto = numero.strip()
    for contacto in agenda:
        if texto in contacto["telefono"]:
            encontrados.append(contacto)
    return encontrados


def eliminar_contacto(agenda, nombre):
    """Elimina el primer contacto con ese nombre exacto.

    No distingue mayusculas ni espacios.
    Devuelve True si lo borra y False si no lo encuentra.
    """
    texto = nombre.strip().lower()
    for i in range(len(agenda)):
        if agenda[i]["nombre"].lower() == texto:
            del agenda[i]
            return True
    return False


def ordenar_contactos(agenda, descendente=False):
    """Devuelve una lista nueva con los contactos ordenados por nombre.

    No distingue mayusculas. Si descendente es True, ordena de Z a A.
    La lista original no se modifica.
    """
    return sorted(agenda, key=lambda c: c["nombre"].lower(), reverse=descendente)


def mostrar_contactos(agenda, descendente=False):
    """Muestra todos los contactos de la agenda por pantalla, ordenados.

    Por defecto orden ascendente (A-Z). Con descendente=True, de Z a A.
    """
    if len(agenda) == 0:
        print("La agenda esta vacia.")
        return
    ordenados = ordenar_contactos(agenda, descendente)
    if descendente:
        print("\nContactos (Z-A):")
    else:
        print("\nContactos (A-Z):")
    for contacto in ordenados:
        print("- " + contacto["nombre"] + ": " + contacto["telefono"])


def mostrar_lista(contactos):
    """Muestra una lista de contactos (para resultados de busqueda)."""
    if len(contactos) == 0:
        print("No se han encontrado contactos.")
        return
    for contacto in contactos:
        print("- " + contacto["nombre"] + ": " + contacto["telefono"])


def guardar_contactos(agenda, ruta=RUTA_DATOS):
    """Guarda la agenda en una hoja de calculo Excel (.xlsx).

    La hoja tiene dos columnas: nombre | telefono.
    La primera fila es la cabecera. El telefono se guarda como texto.
    """
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = NOMBRE_HOJA
    hoja.append([TITULO_NOMBRE, TITULO_TELEFONO])
    for contacto in agenda:
        hoja.append([contacto["nombre"], str(contacto["telefono"])])
    # Formato texto en la columna del telefono (columna B).
    for fila in range(2, hoja.max_row + 1):
        hoja.cell(row=fila, column=2).number_format = "@"
    # Ajuste sencillo del ancho de columnas.
    hoja.column_dimensions["A"].width = 20
    hoja.column_dimensions["B"].width = 18
    libro.save(ruta)


def cargar_contactos(ruta=RUTA_DATOS):
    """Carga la agenda desde una hoja de calculo Excel (.xlsx).

    Si el fichero no existe, devuelve una lista vacia sin dar error.
    Se salta la cabecera y se ignoran las filas vacias o incompletas.
    """
    agenda = []
    try:
        libro = openpyxl.load_workbook(ruta)
    except FileNotFoundError:
        return []
    if NOMBRE_HOJA in libro.sheetnames:
        hoja = libro[NOMBRE_HOJA]
    else:
        hoja = libro.active
    primera = True
    for fila in hoja.iter_rows(values_only=True):
        if primera:
            primera = False
            continue
        if fila is None:
            continue
        if len(fila) < 2:
            continue
        nombre = fila[0]
        telefono = fila[1]
        if nombre is None or telefono is None:
            continue
        nombre = str(nombre).strip()
        telefono = str(telefono).strip()
        if nombre == "" or telefono == "":
            continue
        # Ignorar la cabecera si aparece en otra hoja sin titulo.
        if nombre.lower() == TITULO_NOMBRE and telefono.lower() == TITULO_TELEFONO:
            continue
        agenda.append(crear_contacto(nombre, telefono))
    return agenda


def main():
    """Bucle principal: carga, muestra el menu, actua y guarda al salir."""
    agenda = cargar_contactos()

    seguir = True
    while seguir:
        mostrar_menu()
        opcion = pedir_opcion()

        if opcion == "1":
            nombre = pedir_nombre()
            telefono = pedir_telefono()
            if nombre == "" or telefono == "":
                print("El nombre y el telefono no pueden estar vacios.")
            else:
                contacto = crear_contacto(nombre, telefono)
                anadir_contacto(agenda, contacto)
                print("Contacto anadido.")
        elif opcion == "2":
            nombre = pedir_nombre()
            resultados = buscar_por_nombre(agenda, nombre)
            mostrar_lista(resultados)
        elif opcion == "3":
            numero = pedir_telefono()
            resultados = buscar_por_numero(agenda, numero)
            mostrar_lista(resultados)
        elif opcion == "4":
            nombre = pedir_nombre()
            if eliminar_contacto(agenda, nombre):
                print("Contacto eliminado.")
            else:
                print("No se ha encontrado ese contacto.")
        elif opcion == "5":
            descendente = pedir_orden()
            mostrar_contactos(agenda, descendente)
        elif opcion == "0":
            guardar_contactos(agenda)
            print("Agenda guardada. Hasta pronto.")
            seguir = False
        else:
            print("Opcion no valida. Elige un numero del 0 al 5.")


if __name__ == "__main__":
    main()

